"""Small executable report pipeline, optional native-format dependencies.

Produces editable Office files and a PDF from validated structured rows. Not a
replacement for the richer design skills or an assurance of rendered layout.
"""
import argparse
from html import escape
import json
import math
from pathlib import Path
import re
import tempfile
import os


def build_documents(output,title,rows):
    if not isinstance(title,str) or not title.strip() or len(title)>120:
        raise ValueError('title must be 1..120 characters')
    def text_ok(value):
        return isinstance(value,str) and 0<len(value)<=120 and not re.search(r'[\x00-\x1f\x7f-\x9f\u202a-\u202e]',value)
    if not text_ok(title) or not isinstance(rows,list) or not 1<=len(rows)<=100:
        raise ValueError('1..100 rows required, no control characters')
    for row in rows:
        if not isinstance(row,dict) or not text_ok(row.get('name')) or type(row.get('value')) not in (int,float) or not math.isfinite(row['value']) or abs(row['value'])>1e100:
            raise ValueError('each row needs a bounded name and finite numeric value')
    output=Path(output)
    if output.exists():raise ValueError('output must be new')
    # Imports precede output creation, so absent dependencies leave no directory.
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent,prefix='.document-build-') as temporary:
        temp=Path(temporary)
        doc=Document();doc.add_heading(title,0)
        table=doc.add_table(rows=1,cols=2);table.style='Light Shading Accent 1'
        table.cell(0,0).text='Name';table.cell(0,1).text='Value'
        for row in rows:
            cells=table.add_row().cells;cells[0].text=row['name'];cells[1].text=str(row['value'])
        doc.save(temp/'report.docx')
        book=Workbook();sheet=book.active;sheet.title='Data';sheet.append(['Name','Value'])
        for row in rows:
            sheet.append([row['name'],row['value']])
            sheet.cell(sheet.max_row,1).data_type='s' # Prevent formula injection.
        sheet.append(['Total',f'=SUM(B2:B{len(rows)+1})'])
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17324D')
        sheet.freeze_panes='A2';sheet.column_dimensions['A'].width=36;sheet.column_dimensions['B'].width=18
        book.save(temp/'report.xlsx');book.close()
        styles=getSampleStyleSheet()
        data=[[Paragraph('Name',styles['BodyText']),Paragraph('Value',styles['BodyText'])]]
        data.extend([[Paragraph(escape(r['name']),styles['BodyText']),str(r['value'])] for r in rows])
        tab=Table(data,colWidths=[350,110],repeatRows=1)
        tab.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E9EFF5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
        SimpleDocTemplate(str(temp/'report.pdf')).build([Paragraph(escape(title),styles['Title']),tab])
        deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        for offset in range(0,len(rows),6):
            slide=deck.slides.add_slide(deck.slide_layouts[5]);slide.shapes.title.text=title
            slide.shapes.title.text_frame.paragraphs[0].font.size=Pt(30)
            frame=slide.shapes.add_textbox(Inches(.7),Inches(1.5),Inches(12),Inches(5.3)).text_frame
            for i,row in enumerate(rows[offset:offset+6]):
                p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
                p.text=f"{row['name']}: {row['value']}";p.font.size=Pt(18)
        deck.save(temp/'report.pptx')
        output.mkdir(mode=0o700)
        for path in temp.iterdir():os.replace(path,output/path.name)
    return {ext:str(output/f'report.{ext}') for ext in ('docx','xlsx','pdf','pptx')}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input',type=Path,required=True,help='JSON with title and rows[{name,value}]')
    p.add_argument('--output',type=Path,required=True)
    args=p.parse_args()
    try:
        data=json.loads(args.input.read_text())
        print(json.dumps(build_documents(args.output,data['title'],data['rows']),indent=2));return 0
    except (ValueError,OSError,KeyError,ImportError,TypeError):
        print('Document build failed; check input and optional dependencies',file=__import__('sys').stderr);return 2


if __name__=='__main__':raise SystemExit(main())
