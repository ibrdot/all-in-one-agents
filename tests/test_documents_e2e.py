"""Actual file creation and readback; no claim of visual/native-office rendering."""
import pytest
from runtime.documents import build_documents


def test_all_four_document_formats_round_trip(tmp_path):
    for module in ('docx','openpyxl','reportlab','pptx','pypdf'):pytest.importorskip(module)
    result=build_documents(tmp_path/'reports','Sales review',[{'name':'North','value':120},{'name':'South','value':80}])
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from pptx import Presentation
    assert Document(result['docx']).tables[0].cell(1,0).text=='North'
    workbook=load_workbook(result['xlsx']);assert workbook.active['B4'].value=='=SUM(B2:B3)';workbook.close()
    assert 'Sales review' in PdfReader(result['pdf']).pages[0].extract_text()
    assert any('Sales review' in shape.text for shape in Presentation(result['pptx']).slides[0].shapes if shape.has_text_frame)


@pytest.mark.parametrize('rows', [[],[{'name':'x','value':float('nan')}],[{'name':'x','value':True}],[{'name':'x','value':'oops'}]])
def test_invalid_document_data_writes_nothing(tmp_path,rows):
    with pytest.raises(ValueError):build_documents(tmp_path/'out','Title',rows)
    assert not (tmp_path/'out').exists()


def test_spreadsheet_labels_are_not_formulas(tmp_path):
    for module in ('docx','openpyxl','reportlab','pptx'):pytest.importorskip(module)
    from openpyxl import load_workbook
    r=build_documents(tmp_path/'out','Report',[{'name':'=WEBSERVICE("https://invalid")','value':1}])
    workbook=load_workbook(r['xlsx']);assert workbook.active['A2'].data_type=='s';workbook.close()


def test_pdf_xml_characters_are_escaped(tmp_path):
    for module in ('docx','openpyxl','reportlab','pptx','pypdf'):pytest.importorskip(module)
    from pypdf import PdfReader
    r=build_documents(tmp_path/'out','A < B & C',[{'name':'<tag> & value','value':2}])
    assert '<tag> & value' in PdfReader(r['pdf']).pages[0].extract_text()
